"""Đánh giá thủ công (#17) và duyệt tập kiểm thử tích hợp (#18) bằng MỘT file Excel cho người dùng điền.

prepare: chạy pipeline trên tin nhắn thật của ViLexNorm test (chọn ngẫu nhiên, seed cố định), lấy các câu có gợi ý
         -> reports/manual_eval/danh_gia_thu_cong.xlsx, gồm 3 sheet: Hướng dẫn, Chấm gợi ý (mỗi dòng một gợi ý,
         cột "Chấp nhận" 1/0 + lý do), Duyệt tập kiểm thử (41 câu draft, cột Quyết định: Giữ/Sửa/Bỏ).
         Không ghi đè file đã có (tránh mất phần người dùng đã chấm) trừ khi có --force.
score:   đọc file đã điền -> Top-1/Top-3 chấp nhận, theo vị trí gợi ý, tách theo bước sửa lỗi đúng/sai, phân bố lý do
         (lý do ghi tự do được gom nhóm theo ly_do_phan_loai.csv) -> outputs/metrics/manual_eval.json,
         và áp quyết định duyệt vào data/test/integration_test.csv (Giữ/Sửa -> status=approved, Bỏ -> xóa dòng).

Chạy:  python -m src.pipeline.manual_eval prepare [--n 40]
       python -m src.pipeline.manual_eval score
"""
import argparse
import json
import random
from pathlib import Path

import pandas as pd

from src.utils.metrics import tokenize

XLSX = Path("reports/manual_eval/danh_gia_thu_cong.xlsx")
INTEGRATION = Path("data/test/integration_test.csv")
REASON_CODES = XLSX.parent / "ly_do_phan_loai.csv"  # Mã câu, Gợi ý số, Lý do của người chấm, Nhóm lý do
SEED = 2026
REASONS = ["Sai/đổi nghĩa", "Thêm hoặc bỏ ý", "Đảo vai (ai làm gì cho ai)", "Không tự nhiên",
           "Còn lỗi chính tả/ngữ pháp", "Khác"]
DECISIONS = ["Giữ", "Sửa", "Bỏ"]
GROUPS = {"A": "A – lỗi dấu", "B": "B – viết tắt/không dấu", "C": "C – cần diễn đạt lại", "D": "D – câu đã đúng"}

GUIDE = [
    ("ĐÁNH GIÁ THỦ CÔNG", ""),
    ("Mất khoảng", "45–60 phút cho sheet 'Chấm gợi ý', 10–15 phút cho sheet 'Duyệt tập kiểm thử'. Có thể làm nhiều lần, "
                   "nhớ bấm Lưu (Ctrl+S) và đóng file trước khi báo Claude tính điểm."),
    ("", ""),
    ("SHEET 'Chấm gợi ý'", "Mỗi dòng là một gợi ý diễn đạt lại mà hệ thống đưa ra cho một tin nhắn thật (lấy từ ViLexNorm)."),
    ("Cột 'Ý người viết'", "là câu chuẩn do người gán nhãn viết = điều người gửi thật sự muốn nói. Chấm gợi ý so với cột này."),
    ("Cột 'Chấp nhận'", "chọn 1 nếu gợi ý đạt CẢ BA điều: (1) đúng ý người viết, (2) tự nhiên, bạn sẵn sàng gửi câu đó, "
                        "(3) không thêm hoặc bỏ thông tin. Thiếu một điều thì chọn 0."),
    ("Cột 'Lý do (nếu 0)'", "chọn lý do chính khiến gợi ý không đạt (chọn trong danh sách). Không bắt buộc nhưng giúp phân tích lỗi."),
    ("Ví dụ chấm 1", "Ý: 'Mai mình đi học lúc 7h nhé'  ->  'Ngày mai 7 giờ mình đi học nhé'"),
    ("Ví dụ chấm 0", "Ý: 'Em gửi chị file rồi ạ'  ->  'Chị đã gửi file cho em rồi' (Đảo vai)"),
    ("Ví dụ chấm 0", "Ý: 'Đừng quên mua sữa nhé'  ->  'Nhớ mua sữa và bánh mì nhé' (Thêm hoặc bỏ ý)"),
    ("Lưu ý", "Không cần so các gợi ý với nhau; chấm từng dòng độc lập. Gặp câu khó quyết định thì chọn theo cảm nhận đầu tiên."),
    ("", ""),
    ("SHEET 'Duyệt tập kiểm thử'", "41 câu kiểm thử do AI soạn (chưa ai duyệt). Với mỗi câu, xem câu đúng và các mục phải giữ có hợp lý không."),
    ("Cột 'Quyết định'", "Giữ = câu đúng như ghi; Sửa = điền cách sửa vào hai cột bên phải; Bỏ = câu không hợp lý, loại khỏi tập."),
    ("Cột 'Phải giữ'", "các số, tên riêng, phủ định mà hệ thống không được làm mất, ngăn cách bằng dấu ;"),
    ("Nhóm", "A – lỗi dấu; B – viết tắt/không dấu; C – câu đúng nhưng cần diễn đạt lại (câu đúng = chính đầu vào); "
             "D – câu đã đúng, hệ thống không được sửa."),
]


def sample_messages(data: str, n_pool: int) -> pd.DataFrame:
    df = pd.read_json(data, lines=True)
    df = df[df.target.map(lambda s: 5 <= len(tokenize(s)) <= 30)].drop_duplicates("target")
    return df.sample(min(n_pool, len(df)), random_state=SEED).reset_index(drop=True)


def rating_rows(messages: pd.DataFrame, results: list[dict], n: int) -> tuple[pd.DataFrame, dict]:
    """Một dòng mỗi gợi ý, cho n tin nhắn đầu tiên có gợi ý; thứ tự tin nhắn xáo ngẫu nhiên."""
    with_sugg = [(m, r) for m, r in zip(messages.itertuples(), results) if r["suggestions"]]
    chosen = with_sugg[:n]
    random.Random(SEED).shuffle(chosen)
    rows = []
    for i, (m, r) in enumerate(chosen, 1):
        for k, s in enumerate(r["suggestions"], 1):
            rows.append({"Mã câu": f"M{i:02d}", "Tin nhắn gốc": m.source, "Ý người viết": m.target,
                         "Gợi ý số": k, "Gợi ý": s["text"], "Chấp nhận": None, "Lý do (nếu 0)": None, "Ghi chú": None,
                         "_corrected": r["corrected"]})
    meta = {"pool": len(messages), "with_suggestions": len(with_sugg), "rated_messages": len(chosen),
            "correction_exact": sum(r["corrected"] == m.target for m, r in chosen)}
    return pd.DataFrame(rows), meta


def score_ratings(df: pd.DataFrame) -> dict:
    """df: các dòng 'Chấm gợi ý' đã điền. Tin nhắn chưa chấm đủ mọi gợi ý thì bỏ qua (báo trong n_incomplete)."""
    df = df.assign(ok=pd.to_numeric(df["Chấp nhận"], errors="coerce"))
    done = df.groupby("Mã câu").ok.transform(lambda s: s.notna().all())
    rated = df[done]
    per_msg = rated.sort_values("Gợi ý số").groupby("Mã câu").ok
    n = rated["Mã câu"].nunique()
    reasons = rated[rated.ok == 0]["Lý do (nếu 0)"].fillna("(không ghi)").value_counts()
    return {
        "n_messages": n,
        "n_suggestions": len(rated),
        "n_incomplete": int(df["Mã câu"].nunique() - n),
        "top1_acceptance": float(per_msg.first().mean()) if n else None,
        "top3_acceptance": float((per_msg.max() == 1).mean()) if n else None,
        "suggestion_acceptance": float(rated.ok.mean()) if n else None,
        "acceptance_by_position": {int(k): float(v) for k, v in rated.groupby("Gợi ý số").ok.mean().items()},
        "reasons": {k: int(v) for k, v in reasons.items()},
    }


def apply_review(cases: pd.DataFrame, review: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """Áp quyết định duyệt lên tập kiểm thử; dòng chưa có quyết định giữ nguyên status=draft."""
    cases = cases.copy()
    counts = {d: 0 for d in DECISIONS} | {"chưa duyệt": 0}
    drop = []
    for r in review.itertuples(index=False):
        idx = cases.index[cases.id == r.id]
        decision = r.decision.strip().capitalize() if isinstance(r.decision, str) else ""  # "giữ" -> "Giữ"
        if decision not in DECISIONS:
            counts["chưa duyệt"] += 1
            continue
        counts[decision] += 1
        if len(idx) == 0:  # đã bỏ ở lần chạy trước
            continue
        if decision == "Bỏ":
            drop += list(idx)
            continue
        if decision == "Sửa":
            if isinstance(r.new_correction, str) and r.new_correction.strip():
                cases.loc[idx, "expected_correction"] = r.new_correction.strip()
            if isinstance(r.new_must_keep, str):
                cases.loc[idx, "must_keep"] = r.new_must_keep.strip()
        cases.loc[idx, "status"] = "approved"
    return cases.drop(index=drop), counts


def write_xlsx(path: Path, ratings: pd.DataFrame, cases: pd.DataFrame, meta: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    head_fill, head_font = PatternFill("solid", fgColor="17365D"), Font(bold=True, color="FFFFFF")
    band = PatternFill("solid", fgColor="EAF2F8")
    todo = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")
    wb = Workbook()

    ws = wb.active
    ws.title = "Hướng dẫn"
    for a, b in GUIDE:
        ws.append([a, b])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 28, 110
    for row in ws.iter_rows():
        for c in row:
            c.alignment = wrap

    def sheet(title, df, widths, input_cols, validations):
        ws = wb.create_sheet(title)
        ws.append(list(df.columns))
        for c in ws[1]:
            c.fill, c.font, c.alignment = head_fill, head_font, wrap
        prev, shade = None, False
        for row in df.itertuples(index=False):
            ws.append([None if pd.isna(v) else v for v in row])
            if row[0] != prev:
                shade, prev = not shade, row[0]
            for j, c in enumerate(ws[ws.max_row], 1):
                c.alignment = wrap
                if df.columns[j - 1] in input_cols:
                    c.fill = todo
                elif shade:
                    c.fill = band
        for col, w in zip("ABCDEFGHIJ", widths):
            ws.column_dimensions[col].width = w
        for col_name, options in validations.items():
            col = "ABCDEFGHIJ"[list(df.columns).index(col_name)]
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{len(df) + 1}")
        ws.freeze_panes = "A2"
        return ws

    sheet("Chấm gợi ý", ratings.drop(columns=["_corrected"]), [8, 40, 40, 8, 45, 11, 24, 25],
          {"Chấp nhận", "Lý do (nếu 0)", "Ghi chú"}, {"Chấp nhận": ["1", "0"], "Lý do (nếu 0)": REASONS})
    review = pd.DataFrame({"Mã": cases.id, "Nhóm": cases.group.map(GROUPS), "Đầu vào": cases.input,
                           "Câu đúng": cases.expected_correction, "Phải giữ": cases.must_keep, "Ghi chú AI": cases.notes,
                           "Quyết định": None, "Sửa câu đúng thành": None, "Sửa 'Phải giữ' thành": None})
    sheet("Duyệt tập kiểm thử", review, [6, 20, 36, 36, 14, 24, 11, 32, 18],
          {"Quyết định", "Sửa câu đúng thành", "Sửa 'Phải giữ' thành"}, {"Quyết định": DECISIONS})

    hidden = wb.create_sheet("_meta")
    hidden.append(["meta", json.dumps(meta, ensure_ascii=False)])
    for i, c in enumerate(ratings["_corrected"], 2):
        hidden.cell(i, 1, c)
    hidden.sheet_state = "hidden"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def prepare(args):
    if XLSX.exists() and not args.force:
        raise SystemExit(f"{XLSX} đã có (có thể đã chấm). Dùng --force nếu thật sự muốn tạo lại.")
    from src.pipeline.assistant import from_config

    messages = sample_messages(args.data, args.pool)
    results = from_config(args.config).run_batch(messages.source.tolist())
    ratings, meta = rating_rows(messages, results, args.n)
    meta |= {"data": args.data, "seed": SEED}
    cases = pd.read_csv(INTEGRATION, dtype=str, keep_default_na=False)
    write_xlsx(XLSX, ratings, cases[cases.status == "draft"], meta)
    print(f"-> {XLSX}: {meta['rated_messages']} tin nhắn, {len(ratings)} gợi ý cần chấm; "
          f"{(cases.status == 'draft').sum()} câu kiểm thử cần duyệt")
    print(json.dumps(meta, ensure_ascii=False))


def score(args):
    ratings = pd.read_excel(XLSX, sheet_name="Chấm gợi ý")
    hidden = pd.read_excel(XLSX, sheet_name="_meta", header=None)
    meta = json.loads(hidden.iloc[0, 1])
    ratings["_corrected"] = hidden.iloc[1:len(ratings) + 1, 0].tolist()
    if REASON_CODES.exists():  # lý do ghi tự do -> nhóm lý do (phân loại thủ công, xem file)
        codes = pd.read_csv(REASON_CODES).set_index(["Mã câu", "Gợi ý số"])["Nhóm lý do"]
        keys = list(zip(ratings["Mã câu"], ratings["Gợi ý số"]))
        ratings["Lý do (nếu 0)"] = [codes.get(k, r) for k, r in zip(keys, ratings["Lý do (nếu 0)"])]
    correct = ratings._corrected == ratings["Ý người viết"]
    report = {"meta": meta, "ratings": score_ratings(ratings),
              "correction_exact": score_ratings(ratings[correct]),
              "correction_wrong": score_ratings(ratings[~correct])}

    review = pd.read_excel(XLSX, sheet_name="Duyệt tập kiểm thử", dtype=str)
    review = review.rename(columns={"Mã": "id", "Quyết định": "decision", "Sửa câu đúng thành": "new_correction",
                                    "Sửa 'Phải giữ' thành": "new_must_keep"})
    cases = pd.read_csv(INTEGRATION, dtype=str, keep_default_na=False)
    new_cases, counts = apply_review(cases, review)
    report["review"] = counts
    if any(counts[d] for d in DECISIONS):
        new_cases.to_csv(INTEGRATION, index=False, encoding="utf-8")

    Path("outputs/metrics/manual_eval.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("prepare")
    p.add_argument("--data", default="data/processed/vilexnorm/test.jsonl")
    p.add_argument("--n", type=int, default=40, help="số tin nhắn cần chấm (mỗi tin ~2-3 gợi ý)")
    p.add_argument("--pool", type=int, default=120, help="số tin nhắn lấy mẫu để chạy pipeline")
    p.add_argument("--config", default="configs/pipeline.yaml")
    p.add_argument("--force", action="store_true")
    sub.add_parser("score")
    args = ap.parse_args()
    {"prepare": prepare, "score": score}[args.cmd](args)


if __name__ == "__main__":
    main()
